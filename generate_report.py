import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_image_safe(doc, path, width=Inches(6.0)):
    if os.path.exists(path):
        doc.add_picture(path, width=width)
    else:
        doc.add_paragraph(f"(Không tìm thấy ảnh: {path})").font.color.rgb = RGBColor(255, 0, 0)

def main():
    doc = Document()

    # Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # Title
    title = doc.add_heading('BÁO CÁO CÁC TÍNH NĂNG CHI TIẾT - MINDX GAME ASSET GENERATOR', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('Báo cáo phân tích chi tiết các tính năng đã được hiện thực và kiểm thử thành công trên hệ thống, dựa vào dữ liệu chụp màn hình (Test-sources).')

    # Feature 1
    doc.add_heading('1. Sinh Sprite Nhân Vật (Text to Sprite)', level=1)
    doc.add_paragraph('Cho phép nhập mô tả (prompt) để sinh nhân vật. Cấu hình: Loại Asset (Nhân vật/Sprite), Phong cách (Chibi, Anime, Pixel...), Góc nhìn (Chính diện, Side...). Hỗ trợ tự động tách nền (Transparent Background).')
    add_image_safe(doc, r'Test-sources\Screenshot (1810).png')

    # Feature 2
    doc.add_heading('2. Sinh Cảnh Quan (Text to Background)', level=1)
    doc.add_paragraph('Tính năng sinh hình nền game chi tiết. Hỗ trợ chọn thời gian trong ngày (Ban đêm/Ban ngày), Tỉ lệ (16:9, v.v...) và phong cách (Realistic, 3D Render).')
    add_image_safe(doc, r'Test-sources\Screenshot (1811).png')

    # Feature 3
    doc.add_heading('3. Sinh Gạch Môi Trường (Text to Tileset)', level=1)
    doc.add_paragraph('Sinh các khối gạch môi trường ghép nối được. Hỗ trợ góc nhìn Isometric (2.5D) chuẩn xác. Tích hợp sẵn bộ công cụ cắt gạch (Slicer) tải về dạng file nén ZIP.')
    add_image_safe(doc, r'Test-sources\Screenshot (1812).png')

    # Feature 4
    doc.add_heading('4. Sinh Hoạt Ảnh (Text to Sprite Sheet / Animation)', level=1)
    doc.add_paragraph('Tạo nguyên một dải hành động cho nhân vật (VD: Hành động Running, 3 khung hình). Hỗ trợ xem trước hoạt ảnh trực tiếp (Animation Preview) với tốc độ FPS tuỳ chỉnh trước khi cắt.')
    add_image_safe(doc, r'Test-sources\Screenshot (1814).png')

    # Feature 5
    doc.add_heading('5. Thiết Kế Pixel Art Chuyên Dụng (Text to Pixel)', level=1)
    doc.add_paragraph('Chế độ sinh ảnh Pixel thuần tuý với các thông số chuyên sâu của hệ màu Retro: Cấu hình kích thước khối (Block Size), Giới hạn số lượng màu (Ví dụ: 32 màu).')
    add_image_safe(doc, r'Test-sources\Screenshot (1816).png')

    # Feature 6
    doc.add_heading('6. Chuyển Đổi Từ Ảnh Gốc Sang Sprite (Image to Sprite)', level=1)
    doc.add_paragraph('Cho phép tải lên ảnh phác thảo/ảnh thật, từ đó AI sẽ vẽ lại theo phong cách mới (Anime, Chibi...) với góc nhìn định sẵn. Tự động tách nền cho kết quả.')
    add_image_safe(doc, r'Test-sources\Screenshot (1817).png')

    # Feature 7
    doc.add_heading('7. Chuyển Đổi Từ Ảnh Gốc Sang Pixel (Image to Pixel Art)', level=1)
    doc.add_paragraph('Tính năng tải ảnh lên và chuyển đổi (filter hóa) trực tiếp thành dạng Pixel Art chuẩn xác dùng trong game.')
    add_image_safe(doc, r'Test-sources\Screenshot (1818).png')

    # Feature 8
    doc.add_heading('8. Thư Viện Asset (Asset Library)', level=1)
    doc.add_paragraph('Trang quản lý các Asset đã sinh ra. Hiển thị thông tin siêu dữ liệu (metadata) bao gồm mô tả, phong cách, và nhãn phân loại (TILESET, SPRITE, PIXEL_ART, BACKGROUND...).')
    add_image_safe(doc, r'Test-sources\image copy.png')

    # Video Note
    doc.add_heading('9. Video Demo Toàn Cảnh', level=1)
    doc.add_paragraph('Quá trình thao tác UI thực tế, tốc độ sinh ảnh, và trải nghiệm cắt khung hình động (Animation) được ghi lại chi tiết trong file video: Test-sources\\2026-08-06 20-15-41.mp4')
    
    # Save the document
    report_path = 'Bao_Cao_Tinh_Nang_Chi_Tiet_MindXAsset.docx'
    doc.save(report_path)
    print(f"Report successfully saved to {report_path}")

if __name__ == '__main__':
    main()
